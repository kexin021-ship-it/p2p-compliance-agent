from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"C:\Users\kesin\.codex\plugins\cache\openai-curated-remote"
    r"\openai-templates\0.1.1\skills\artifact-template-experiment-analysis"
    r"\assets\reference.docx"
)
OUTPUT = PROJECT / "docs" / "采购单流程调查Agent对照实验最终报告.docx"
FIGURES = PROJECT / "docs" / "figures"

BLACK = "111111"
MUTED = "555555"
GREEN = "176B3A"
LIGHT_GREEN = "EAF3ED"
LIGHT_GRAY = "F4F4F2"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and keep is not None:
        p_pr.remove(keep)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_body(document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def clear_container(container):
    element = container._element
    for child in list(element):
        element.remove(child)
    paragraph = OxmlElement("w:p")
    element.append(paragraph)


def configure_styles(document):
    normal = document.styles["normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28

    title = document.styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)

    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Microsoft YaHei"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)

    for style_name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(13 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_footer(section):
    footer = section.footer
    clear_container(footer)
    table = footer.add_table(rows=1, cols=2, width=Inches(6.9))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.9)
    table.columns[1].width = Inches(1.0)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.text = "采购单流程调查 Agent 对照实验"
    right.text = ""
    for run in left.paragraphs[0].runs:
        set_run_font(run, size=8, color=GREEN)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right.paragraphs[0].add_run()
    set_run_font(run, size=8, color=BLACK)
    add_field(run, "PAGE")
    set_table_borders(table, color=WHITE, size="0")


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(paragraph)
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="normal")
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.16)
    run = paragraph.add_run("•  " + text)
    set_run_font(run)
    return paragraph


def add_numbered(document, number, text):
    paragraph = document.add_paragraph(style="normal")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.28)
    run = paragraph.add_run(f"{number}  {text}")
    set_run_font(run)
    return paragraph


def add_table(document, headers, rows, widths=None, header_fill=GREEN):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = str(value)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, size=9, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=8.6)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(document, filename, caption, width=6.55):
    path = FIGURES / filename
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    paragraph.paragraph_format.space_after = Pt(3)
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    run = cap.add_run(caption)
    set_run_font(run, size=8.5, color=MUTED)


def page_break(document):
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_document():
    document = Document(str(TEMPLATE))
    clear_body(document)
    configure_styles(document)
    section = document.sections[0]
    configure_footer(section)

    # Cover page preserves the template's restrained, spacious report composition.
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(18)
    run = kicker.add_run("实验分析报告")
    set_run_font(run, size=12, bold=True, color=GREEN)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("采购单流程调查 Agent\n与规则基线对照实验")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("冻结 V2 2 全量 3000 份采购单实验")

    for _ in range(7):
        document.add_paragraph()

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("项目成果文档\n2026年9月6日")
    set_run_font(run, size=10.5, color=MUTED)
    page_break(document)

    add_heading(document, "执行摘要", 1)
    add_body(
        document,
        "冻结 V2.2 Agent 已在全部 3000 份采购单上稳定完成批处理，五个 Batch 分批均成功，0 个请求失败。Agent 与纯规则基线在 2628 份采购单上决策一致，一致率为 87.6%；在 372 份采购单上存在分歧。",
    )
    add_body(
        document,
        "全量 3000 例没有逐例人工真值，因此 87.6% 只能解释为系统一致率，不能解释为准确率。冻结独立留出集共 47 例，其中双方 9 个分歧全部完成人工裁决；Agent 判对 7 个，规则判对 2 个。Agent 因此比规则净多判对 5 例，对完整 47 例形成精确的相对准确率优势 10.6 个百分点。",
    )
    add_body(
        document,
        "本项目累计人工判断 42 个互不重复的采购单，包括开发集 20 例、初始盲测与回归集 13 例、冻结独立留出分歧集 9 例。只有最后 9 例属于冻结版本后的独立人工裁决，可用于最终无偏比较。",
    )
    add_table(
        document,
        ["关键指标", "结果", "解释"],
        [
            ["全量规模", "3000 PO", "用于输出分布与系统一致性分析"],
            ["规则与 Agent 一致", "2628 PO  87.6%", "不是准确率"],
            ["冻结留出分歧裁决", "9 PO", "规则 2 对  Agent 7 对"],
            ["相对准确率优势", "+10.6 个百分点", "Agent 相对规则的精确差值"],
            ["Agent 证据不足", "361 PO  12.0%", "规则基线无法表达的第三类状态"],
        ],
        widths=[1.55, 1.45, 3.65],
    )

    add_heading(document, "一 研究问题与实验假设", 1)
    add_body(
        document,
        "研究问题是，在采购到付款流程日志的调查任务中，基于大语言模型的 Agent 是否能比纯规则基线提供更可靠的采购单级决策，同时明确指出重点行项目、主要发现与证据缺口。",
    )
    add_body(
        document,
        "实验假设：在冻结提示词、事实摘要逻辑和模型版本后，当 Agent 与规则基线产生分歧时，Agent 在独立人工盲审中的正确次数将高于规则；其主要增量应来自对证据不足状态的显式识别，而不是简单扩大或缩小调查范围。",
    )

    add_heading(document, "二 系统与决策定义", 1)
    add_heading(document, "规则基线", 2)
    add_body(
        document,
        "规则基线完全不调用 AI，依据预先定义的事件顺序、必需事件、撤销与闭环规则，输出需要进一步调查或无需进一步调查。它提供了项目所需的笨办法对照组。",
    )
    add_heading(document, "冻结 Agent", 2)
    add_body(
        document,
        "冻结 V2.2 Agent 使用 gpt-5.4-mini-2026-03-17，reasoning effort 为 low。输入为结构化采购单事实摘要，输出采购单级决策、主要发现、重点行项目、观察事实、证据、可能原因、缺失信息、决策理由与建议动作。",
    )
    add_table(
        document,
        ["决策", "中文含义", "适用条件"],
        [
            ["further_investigation", "需要进一步调查", "存在明确、尚未解决的流程异常或缺口"],
            ["insufficient_evidence", "证据不足", "日志缺少关键状态值或后续事件，无法可靠定性"],
            ["no_further_investigation", "无需进一步调查", "事件显示流程闭环、撤回后终止或无未解决信号"],
        ],
        widths=[1.8, 1.5, 3.35],
    )

    add_heading(document, "三 样本设计与人工标注", 1)
    add_body(
        document,
        "实验采用开发、回归、冻结留出和全量运行四个层次。开发集与初始盲测结果参与了 V2.2 的形成，因此仅用于调试与稳健性检查；冻结独立留出集在版本锁定后运行，用于最终相对比较。",
    )
    add_table(
        document,
        ["阶段", "样本", "人工判断", "用途", "最终无偏比较"],
        [
            ["开发集", "20 PO", "20", "提示词与输出逻辑调优", "否"],
            ["初始盲测", "60 PO", "13", "错误发现与 V2.2 回归", "否"],
            ["冻结独立留出", "47 PO", "9", "裁决双方全部分歧", "是"],
            ["冻结全量运行", "3000 PO", "无逐例真值", "规模与分流分析", "不适用"],
        ],
        widths=[1.25, 1.05, 1.05, 2.3, 1.0],
    )
    add_body(
        document,
        "人工判断总量为 42 个互不重复的采购单，而不是 20 加 9。对应标签文件为 development_labels_v1.jsonl 共 20 条、blind_labels_v1.jsonl 共 13 条、blind_holdout_labels_v2_2.jsonl 共 9 条。development_labels.jsonl 与 development_labels_v1.jsonl 内容及 SHA256 相同，是同一组标签的重复命名，不应重复计数。",
    )

    add_heading(document, "四 冻结与复现控制", 1)
    add_bullet(document, "冻结版本指纹  6D016275F48316DF74CDCB2A3336D94DEFEA7FC43B74B479000AFD2678B3DC6B")
    add_bullet(document, "全量预测 SHA256  2D2D85D72D0EA45F1CD495406E87B1556FD7F61A7C028774AAFEDF6079B38584")
    add_bullet(document, "全量预测数量  3000 个唯一采购单")
    add_bullet(document, "批处理状态  5 个分批全部完成  失败请求 0")
    add_body(
        document,
        "版本指纹用于确认模型、系统提示词和摘要程序未在冻结后发生变化；预测文件哈希用于确认后续评分、汇总和文档引用的是同一份全量结果。",
    )

    page_break(document)
    add_heading(document, "五 全量 3000 份采购单结果", 1)
    add_heading(document, "决策分布", 2)
    add_figure(
        document,
        "full_3000_agent_distribution_v2_2.png",
        "图一  冻结 V2.2 Agent 在 3000 份采购单上的决策分布",
    )
    add_table(
        document,
        ["系统", "需要进一步调查", "证据不足", "无需进一步调查"],
        [
            ["规则基线", "785  26.2%", "0", "2215  73.8%"],
            ["冻结 V2.2 Agent", "665  22.2%", "361  12.0%", "1974  65.8%"],
        ],
        widths=[1.55, 1.7, 1.45, 1.8],
    )
    add_body(
        document,
        "Agent 将二元规则分流扩展为三类决策。361 份采购单被明确标记为证据不足，避免把日志状态缺失直接解释成业务异常或流程闭环。",
    )

    add_heading(document, "系统一致性与分歧", 2)
    add_figure(
        document,
        "full_3000_disagreement_breakdown_v2_2.png",
        "图二  372 个规则与 Agent 分歧的构成",
    )
    add_table(
        document,
        ["规则决策", "Agent 决策", "PO 数量"],
        [
            ["需要进一步调查", "需要进一步调查", "665"],
            ["需要进一步调查", "证据不足", "109"],
            ["需要进一步调查", "无需进一步调查", "11"],
            ["无需进一步调查", "证据不足", "252"],
            ["无需进一步调查", "无需进一步调查", "1963"],
        ],
        widths=[2.35, 2.35, 1.6],
    )
    add_body(
        document,
        "372 个分歧中有 361 个来自 Agent 的证据不足判断，其中 252 个原本被规则判为无需调查，109 个原本被规则判为需要调查。仅 11 个案例被 Agent 从需要调查降为无需调查。由此可见，Agent 的主要行为变化是增加不确定性表达，而非单向放宽或收紧调查标准。",
    )

    add_heading(document, "六 冻结独立留出集比较", 1)
    add_figure(
        document,
        "holdout_disagreement_accuracy_v2_2.png",
        "图三  9 个已裁决分歧案例中的决策正确率",
    )
    add_table(
        document,
        ["指标", "规则基线", "冻结 Agent"],
        [
            ["9 个分歧案例决策正确", "2  22.2%", "7  77.8%"],
            ["主要发现正确", "不适用", "7  77.8%"],
            ["重点行项目正确", "不适用", "9  100.0%"],
            ["对 47 例的相对准确率差", "基准", "+10.6 个百分点"],
        ],
        widths=[3.05, 1.6, 1.75],
    )
    add_body(
        document,
        "47 个独立留出案例中，38 个案例双方预测相同且未人工复核。无论这 38 个共同判断正确与否，Agent 的总正确数量始终比规则多 5 个，因此 5 除以 47 得到的 10.6 个百分点是可确定的相对准确率优势。双方各自的绝对准确率仍无法从现有标签中计算。",
    )
    add_body(
        document,
        "统计解释：9 个裁决分歧的样本量较小。按两侧精确 McNemar 思路，仅考察 7 比 2 的不一致正确数，p 值约为 0.18，未达到常用的 0.05 显著性阈值。因此结果支持 Agent 更优的方向性判断，但不能视为强统计定论。",
    )

    add_heading(document, "七 主要发现与业务解释", 1)
    add_table(
        document,
        ["Agent 主要发现", "PO 数量", "占比"],
        [
            ["流程看起来已解决", "1974", "65.8%"],
            ["流程状态不确定", "360", "12.0%"],
            ["发票入账后未观察到清账", "335", "11.2%"],
            ["未观察到必需的发票入账", "175", "5.8%"],
            ["未观察到必需的收货", "55", "1.8%"],
            ["供应商开票早于采购单创建", "51", "1.7%"],
            ["其他撤销与顺序问题", "50", "1.7%"],
        ],
        widths=[3.85, 1.25, 1.15],
    )
    add_body(
        document,
        "规则适合识别定义清楚、模式稳定的异常，成本低且输出一致；Agent 的额外价值集中在跨行项目整合、因果解释和不确定性表达。实际部署时，证据不足应进入补数或复核队列，而不是与明确异常混在同一调查队列中。",
    )

    page_break(document)
    add_heading(document, "八 运行成本与工程可靠性", 1)
    add_table(
        document,
        ["项目", "数量"],
        [
            ["普通输入 Token", "4,156,666"],
            ["缓存输入 Token", "8,219,904"],
            ["输出 Token  含推理", "1,981,847"],
            ["总 Token", "14,358,417"],
            ["估算 Batch API 费用", "约 12.65 美元"],
        ],
        widths=[3.45, 2.75],
    )
    add_body(
        document,
        "费用根据运行记录中的 Token 使用量和报告生成时记录的 GPT-5.4 Mini Batch 单价估算，最终金额以 OpenAI API 账单为准。五个批次均完成且无失败请求，说明冻结版本具备 3000 份采购单规模的稳定批处理能力。",
    )

    add_heading(document, "九 局限性与有效性边界", 1)
    add_numbered(document, "1", "全量 3000 例没有逐例人工真值，一致率不能被解释为准确率。")
    add_numbered(document, "2", "独立留出集只人工裁决 9 个分歧，38 个共同预测未复核，因此绝对准确率不可得。")
    add_numbered(document, "3", "已裁决分歧只有 9 例，统计功效有限，10.6 个百分点应表述为精确的相对差值与方向性证据。")
    add_numbered(document, "4", "开发集 20 例和初始盲测人工复核 13 例参与过 V2.2 调优，不能重新包装为独立测试结果。")
    add_numbered(document, "5", "事件日志不包含所有金额、数量、状态变更后的值、付款凭证和日志截断后的处理，因此业务事实结论仍受数据范围限制。")
    add_numbered(document, "6", "本实验验证的是给定日志与规则定义下的流程调查能力，不直接等价于财务审计结论或真实损失识别能力。")

    add_heading(document, "十 结论与下一步", 1)
    add_body(
        document,
        "结论：现有独立证据更支持冻结 V2.2 Agent，而不是纯规则基线。Agent 在 47 个独立留出案例上形成 10.6 个百分点的相对准确率优势，并在 3000 份采购单中把 12.0% 的案例识别为证据不足。项目已完成从规则基线、Agent 开发、冻结留出验证到 3000 份采购单全量运行的主要实验闭环。",
    )
    add_body(
        document,
        "下一步不必再修改冻结 V2.2。若需要把项目升级为更强的论文证据，应优先补充独立人工标签，而不是继续调提示词。最有价值的扩展是从 38 个共同预测中随机抽取审计样本，或建立更大的全新留出集，以估计双方的绝对准确率和置信区间。",
    )

    add_heading(document, "附录一 人工标签文件", 1)
    add_table(
        document,
        ["标签文件", "记录数", "角色"],
        [
            [r"evals\development_labels_v1.jsonl", "20", "开发集人工答案"],
            [r"evals\development_labels.jsonl", "20", "与上一文件完全相同的重复命名"],
            [r"evals\blind_labels_v1.jsonl", "13", "初始盲测与回归人工答案"],
            [r"evals\blind_holdout_labels_v2_2.jsonl", "9", "冻结独立留出分歧人工答案"],
        ],
        widths=[3.55, 0.85, 2.1],
    )

    add_heading(document, "附录二 可复现文件", 1)
    add_table(
        document,
        ["文件", "用途"],
        [
            [r"outputs\agent_predictions_all_3000_frozen_v2_2.jsonl", "全量预测"],
            [r"outputs\agent_predictions_all_3000_frozen_v2_2_summary.json", "全量原始汇总"],
            [r"outputs\final_experiment_metrics_v2_2.json", "报告统计"],
            [r"outputs\full_3000_agent_distribution_v2_2.csv", "决策分布"],
            [r"outputs\full_3000_rule_agent_crosstab_v2_2.csv", "交叉表"],
            [r"outputs\final_holdout_metrics_v2_2.csv", "留出集指标"],
        ],
        widths=[5.15, 1.25],
    )

    add_heading(document, "建议引用表述", 2)
    quote = (
        "在 3000 份采购单上的全量实验中，冻结 V2.2 Agent 与纯规则基线达到 87.6% 的一致率，并将 12.0% 的采购单识别为证据不足。在 47 份冻结独立留出案例中，双方 9 个分歧均经人工裁决，Agent 判对 7 个、规则判对 2 个，因此 Agent 相对规则净多判对 5 份采购单，对完整留出集形成 10.6 个百分点的相对准确率优势。由于其余共同预测案例未全部获得人工标签，本研究不将全量一致率解释为绝对准确率。"
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.right_indent = Inches(0.35)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(10)
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepLines"))
    run = paragraph.add_run(quote)
    set_run_font(run, size=10, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
